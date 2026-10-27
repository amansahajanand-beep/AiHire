"""Text extraction and structured parsing of jobs and resumes."""

from __future__ import annotations

import io
import re
from datetime import date

from app.models import Job
from app.rag import gemini

_MIN_TEXT_CHARS = 200  # below this a PDF is treated as scanned and sent to OCR


def extract_text(filename: str, data: bytes) -> str:
    lower = (filename or "").lower()
    text = ""
    if lower.endswith(".docx"):
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif lower.endswith(".pdf"):
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception:  # noqa: BLE001 - corrupt PDFs fall through to OCR
            text = ""
        if len(text.strip()) < _MIN_TEXT_CHARS:
            text = gemini.ocr_pdf(data) or text
    else:
        raise ValueError(f"Unsupported resume format: {filename}")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 50:
        raise ValueError("Could not read any text from this resume")
    return text


def _split_list(value: str | None) -> list[str]:
    return [p.strip() for p in re.split(r"[,;\n|•]+", value or "") if p.strip()]


def job_text(job: Job) -> str:
    return "\n".join(
        f"{label}: {value}"
        for label, value in (
            ("Title", job.title),
            ("Department", job.department),
            ("Location", job.location),
            ("Required experience", job.experience),
            ("Mandatory skills", job.skills),
            ("Qualifications", job.qualifications),
            ("Responsibilities", job.responsibilities),
            ("Description", job.description),
        )
        if value and str(value).strip()
    )


def _as_str_list(value) -> list[str]:
    if isinstance(value, str):
        value = _split_list(value)
    return [str(v).strip() for v in (value or []) if str(v).strip()]


def parse_job(job: Job) -> dict:
    """Structured requirements: required_skills[], min_years, qualifications[], responsibilities[], summary."""
    fallback = {
        "required_skills": _split_list(job.skills),
        "min_years": None,
        "qualifications": _split_list(job.qualifications),
        "responsibilities": [
            r.strip(" -•\t") for r in re.split(r"\n+", job.responsibilities or "") if r.strip(" -•\t")
        ],
        "summary": (job.description or "").strip()[:600],
    }
    try:
        data = gemini.generate_json(
            "Parse this job posting into structured requirements. Use only what the posting says; "
            "do not invent requirements. If mandatory skills are not listed, extract them from the description.\n\n"
            'Return JSON: {"required_skills": [short skill/tool names], "min_years": number or null '
            '(minimum years of experience required), "qualifications": [required degrees/certifications], '
            '"responsibilities": [one short sentence each], "summary": "2 sentence role summary"}\n\n'
            f"JOB POSTING\n{job_text(job)}",
            label="job_parse",
        )
    except gemini.GeminiError:
        return fallback
    try:
        min_years = float(data.get("min_years")) if data.get("min_years") is not None else None
    except (TypeError, ValueError):
        min_years = None
    return {
        "required_skills": _as_str_list(data.get("required_skills")) or fallback["required_skills"],
        "min_years": min_years,
        "qualifications": _as_str_list(data.get("qualifications")) or fallback["qualifications"],
        "responsibilities": _as_str_list(data.get("responsibilities")) or fallback["responsibilities"],
        "summary": str(data.get("summary") or fallback["summary"]).strip(),
    }


def parse_resume(text: str, today: date | None = None) -> dict:
    """Structured resume. Extracts only what the resume states (same rules as the original screening prompt)."""
    today = today or date.today()
    data = gemini.generate_json(
        "Extract the following from the candidate resume. Use ONLY information stated in the resume; "
        'never assume. If something is missing return "Not Available" (or an empty list).\n\n'
        "Rules:\n"
        "- skills: only skills, tools, software, technologies, frameworks and role-specific knowledge actually "
        "mentioned or demonstrated in the resume.\n"
        "- experience: every company separately (never merge companies), with designation, start, end and a "
        'short summary of responsibilities. Write dates as "Mon YYYY" or "YYYY"; for a current job write end as '
        f'"Present" (today is {today.isoformat()}).\n'
        "- education: the actual degrees, qualifications and certifications stated.\n"
        "- location: the candidate's current city only (current/present address, else current employer's city, "
        'else last employer\'s city); "Not Available" if not reliable.\n'
        "- name, email and phone exactly as written.\n\n"
        'Return JSON: {"candidate_name": "", "applied_role": "", "email": "", "phone": "", '
        '"location": "", "skills": [""], '
        '"experience": [{"company": "", "designation": "", "start": "", "end": "", "summary": ""}], '
        '"education": [""], "summary": "2-3 sentence factual summary of the profile"}\n\n'
        f"CANDIDATE RESUME\n{text[:30000]}",
        label="resume_parse",
    )
    experience = []
    for e in data.get("experience") or []:
        if isinstance(e, dict):
            experience.append(
                {
                    "company": str(e.get("company") or "").strip(),
                    "designation": str(e.get("designation") or "").strip(),
                    "start": str(e.get("start") or "").strip(),
                    "end": str(e.get("end") or "").strip(),
                    "summary": str(e.get("summary") or "").strip(),
                }
            )
    email = str(data.get("email") or "").strip().lower()
    if not email:
        m = re.search(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", text)
        email = m.group(0).lower() if m else ""
    return {
        "candidate_name": str(data.get("candidate_name") or "").strip(),
        "applied_role": str(data.get("applied_role") or "").strip(),
        "email": email,
        "phone": str(data.get("phone") or "").strip(),
        "location": str(data.get("location") or "").strip(),
        "skills": _as_str_list(data.get("skills")),
        "experience": experience,
        "education": _as_str_list(data.get("education")),
        "summary": str(data.get("summary") or "").strip(),
    }
